from io import BytesIO

import pymupdf
from openpyxl import load_workbook

from app.config import estimate_cost
from app.export import build_workbook, compute_metrics
from app.pdf import is_blank, page_count, render_pages

EXAM = {"subject": "Biology", "exam_date": "2026-09-01", "total_marks": 6}
QUESTIONS = [{"id": "q1", "label": "1", "max_marks": 4}, {"id": "q2", "label": "2", "max_marks": 2}]
SCRIPTS = [{"id": "s2", "roll_number": "21B002"}, {"id": "s1", "roll_number": "21B001"}]


def answer(script, q, ai, final, confidence="green", status="accepted"):
    return {"script_id": script, "question_id": q, "ai_marks": ai, "check_marks": ai,
            "confidence": confidence, "status": status, "final_marks": final,
            "decided_by": "T", "decided_at": "2026-09-02", "comment": None}


ANSWERS = [
    answer("s1", "q1", 3, 3), answer("s1", "q2", 2, 1, "yellow", "changed"),
    answer("s2", "q1", 4, 4), answer("s2", "q2", 1, None, "red", "pending"),
]


def test_workbook_marks_sheet():
    wb = load_workbook(BytesIO(build_workbook(EXAM, QUESTIONS, SCRIPTS, ANSWERS, {"s1": 0.2, "s2": 0.3})))
    assert wb.sheetnames == ["Marks", "Detail", "Pilot metrics"]
    rows = list(wb["Marks"].values)
    assert rows[0] == ("Roll number", "Q1 (/4)", "Q2 (/2)", "Total (/6)", "Status")
    assert rows[1] == ("21B001", 3, 1, 4, "Complete")
    assert rows[2] == ("21B002", 4, None, 4, "1 pending")
    assert wb["Detail"].max_row == 5


def test_metrics():
    m = dict(compute_metrics(ANSWERS, {"s1": 0.2, "s2": 0.3}))
    assert m["Answers decided by teachers"] == 3
    assert m["AI mark accepted unchanged (%)"] == 66.7
    assert m["Mean absolute difference (marks)"] == 0.33
    assert m["Average AI cost per script (USD)"] == 0.25


def test_estimate_cost():
    assert estimate_cost("openai/gpt-oss-120b", 1_000_000, 0) == 0.15
    assert estimate_cost("openai/gpt-oss-120b", 0, 1_000_000) == 0.60
    assert round(estimate_cost("qwen/qwen3.8-27b", 0, 0, cache_read=1_000_000), 6) == 0.4
    assert estimate_cost("unknown/model", 1_000_000, 1_000_000) == 0


def test_render_pages():
    doc = pymupdf.open()
    for i in range(2):
        page = doc.new_page()
        page.insert_text((72, 72), f"Q{i + 1} handwritten answer")
    data = doc.tobytes()
    assert page_count(data) == 2
    images = render_pages(data, dpi=100)
    assert len(images) == 2
    assert all(img.startswith(b"\x89PNG") for img in images)


def test_blank_page_detection():
    doc = pymupdf.open()
    doc.new_page()
    written = doc.new_page()
    written.insert_text((72, 100), "21. Supervised learning uses labelled data", fontsize=24)
    blank_png, written_png = render_pages(doc.tobytes())
    assert is_blank(blank_png)
    assert not is_blank(written_png)
