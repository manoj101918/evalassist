"""Build the exam results workbook: Marks, Detail and Pilot metrics sheets."""
from collections import Counter
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BOLD = Font(bold=True)
FILLS = {
    "green": PatternFill("solid", fgColor="C6EFCE"),
    "yellow": PatternFill("solid", fgColor="FFEB9C"),
    "red": PatternFill("solid", fgColor="FFC7CE"),
}


def _header(ws, values: list[str]) -> None:
    ws.append(values)
    for cell in ws[1]:
        cell.font = BOLD
    ws.freeze_panes = "A2"


def _autosize(ws) -> None:
    for i, col in enumerate(ws.columns, start=1):
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 60)


def compute_metrics(answers: list[dict], cost_by_script: dict[str, float]) -> list[tuple[str, object]]:
    decided = [a for a in answers if a["final_marks"] is not None and a["ai_marks"] is not None]
    n = len(decided)
    diffs = [abs(float(a["final_marks"]) - float(a["ai_marks"])) for a in decided]
    colours = Counter(a["confidence"] for a in answers)
    total_answers = len(answers) or 1
    rows: list[tuple[str, object]] = [
        ("Answers graded by AI", len(answers)),
        ("Answers decided by teachers", n),
        ("AI mark accepted unchanged (%)", round(100 * sum(d == 0 for d in diffs) / n, 1) if n else None),
        ("AI within 0.5 marks of teacher (%)", round(100 * sum(d <= 0.5 for d in diffs) / n, 1) if n else None),
        ("Mean absolute difference (marks)", round(sum(diffs) / n, 2) if n else None),
    ]
    for colour in ("green", "yellow", "red"):
        rows.append((f"{colour.title()} answers (%)", round(100 * colours[colour] / total_answers, 1)))
        in_colour = [abs(float(a["final_marks"]) - float(a["ai_marks"]))
                     for a in decided if a["confidence"] == colour]
        if in_colour:
            rows.append((f"  {colour.title()}: AI accepted unchanged (%)",
                         round(100 * sum(d == 0 for d in in_colour) / len(in_colour), 1)))
    costs = [c for c in cost_by_script.values() if c > 0]
    rows += [
        ("Scripts processed by AI", len(costs)),
        ("Total AI cost (USD)", round(sum(costs), 4)),
        ("Average AI cost per script (USD)", round(sum(costs) / len(costs), 4) if costs else None),
    ]
    return rows


def build_workbook(exam: dict, questions: list[dict], scripts: list[dict],
                   answers: list[dict], cost_by_script: dict[str, float]) -> bytes:
    """
    questions: [{id, label, max_marks}] in display order
    scripts:   [{id, roll_number}]
    answers:   [{script_id, question_id, ai_marks, check_marks, confidence, status,
                 final_marks, decided_by, decided_at, comment}]
    """
    wb = Workbook()
    by_key = {(a["script_id"], a["question_id"]): a for a in answers}

    ws = wb.active
    ws.title = "Marks"
    _header(ws, ["Roll number"] + [f"Q{q['label']} (/{float(q['max_marks']):g})" for q in questions]
            + [f"Total (/{float(exam['total_marks']):g})", "Status"])
    for s in sorted(scripts, key=lambda s: s["roll_number"]):
        row, total, pending = [s["roll_number"]], 0.0, 0
        for q in questions:
            a = by_key.get((s["id"], q["id"]))
            if a is None or a["final_marks"] is None:
                row.append(None)
                pending += 1
            else:
                row.append(float(a["final_marks"]))
                total += float(a["final_marks"])
        row += [total, "Complete" if pending == 0 else f"{pending} pending"]
        ws.append(row)
    _autosize(ws)

    ws = wb.create_sheet("Detail")
    _header(ws, ["Roll number", "Question", "Max marks", "AI marks", "Check marks", "Confidence",
                 "Final marks", "Decision", "Decided by", "Decided at", "Comment"])
    labels = {q["id"]: q for q in questions}
    position = {q["id"]: i for i, q in enumerate(questions)}
    rolls = {s["id"]: s["roll_number"] for s in scripts}
    ordered = sorted(answers, key=lambda a: (rolls.get(a["script_id"], ""),
                                             position.get(a["question_id"], 0)))
    for a in ordered:
        q = labels.get(a["question_id"], {})
        ws.append([rolls.get(a["script_id"]), q.get("label"), float(q.get("max_marks", 0)),
                   a["ai_marks"], a["check_marks"], a["confidence"], a["final_marks"], a["status"],
                   a.get("decided_by"), a.get("decided_at"), a.get("comment")])
        ws.cell(row=ws.max_row, column=6).fill = FILLS.get(a["confidence"], PatternFill())
    _autosize(ws)

    ws = wb.create_sheet("Pilot metrics")
    _header(ws, ["Metric", "Value"])
    ws.append(["Exam", exam["subject"]])
    ws.append(["Exam date", str(exam.get("exam_date") or "")])
    for name, value in compute_metrics(answers, cost_by_script):
        ws.append([name, value])
    _autosize(ws)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
